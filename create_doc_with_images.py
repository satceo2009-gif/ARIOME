from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

doc = Document()

# Title
title = doc.add_heading('AriOme Application', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle = doc.add_paragraph('Screen Documentation with Screenshots')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('')

screens = [
    ('01_welcome.png', '1. Welcome Screen', 'Route: /'),
    ('02_onboarding.png', '2. Onboarding - Welcome', 'Route: /onboarding'),
    ('03_onboarding_intentions.png', '3. Onboarding - Intention Selection', 'Route: /onboarding (Step 2)'),
    ('04_auth_login.png', '4. Authentication - Login', 'Route: /auth'),
    ('05_auth_register.png', '5. Authentication - Register', 'Route: /auth (Register mode)'),
    ('06_reflect.png', '6. Reflect Screen (Self-AriOme)', 'Route: /(tabs)/reflect'),
    ('07_practices.png', '7. Practices Screen', 'Route: /(tabs)/practices'),
    ('08_wisdom.png', '8. Wisdom Library', 'Route: /(tabs)/wisdom'),
    ('09_journal.png', '9. Journal Screen', 'Route: /(tabs)/journal'),
    ('10_circles.png', '10. Circles (Community)', 'Route: /(tabs)/circles'),
    ('11_profile.png', '11. Profile Screen', 'Route: /(tabs)/profile'),
]

for filename, title_text, route in screens:
    filepath = f'/app/screenshots/{filename}'
    
    # Add title
    heading = doc.add_heading(title_text, level=1)
    
    # Add route
    route_para = doc.add_paragraph(route)
    route_para.runs[0].italic = True
    
    # Add image if exists
    if os.path.exists(filepath):
        doc.add_picture(filepath, width=Inches(6.5))
        last_paragraph = doc.paragraphs[-1]
        last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        doc.add_paragraph(f'[Screenshot not available: {filename}]')
    
    # Add spacing
    doc.add_paragraph('')

# Footer
doc.add_paragraph('')
doc.add_paragraph('─' * 50)
doc.add_paragraph('Document generated: January 2025')
doc.add_paragraph('AriOme - Your Path to Inner Peace')

# Save
doc.save('/app/AriOme_Screens_With_Screenshots.docx')
print('Document saved: /app/AriOme_Screens_With_Screenshots.docx')

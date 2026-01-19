from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import os

# Create document
doc = Document()

# Set up styles
style = doc.styles['Title']
style.font.size = Pt(28)
style.font.bold = True

style = doc.styles['Heading 1']
style.font.size = Pt(18)
style.font.bold = True

style = doc.styles['Heading 2']
style.font.size = Pt(14)
style.font.bold = True

# Title
title = doc.add_heading('AriOme Application', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle = doc.add_paragraph('Screen Documentation')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('A conscious-living mobile ecosystem for self-evolution and mindful reflection')
doc.add_paragraph('')

# Screens
screens = [
    {
        "title": "1. Welcome Screen",
        "route": "/",
        "description": """• Landing page introducing AriOme
• Features the AriOme logo with lotus/infinity symbol
• Tagline: "Your Path to Inner Peace"
• "Begin Your Journey" call-to-action button
• Dark, calming background with teal accent colors"""
    },
    {
        "title": "2. Onboarding - Welcome",
        "route": "/onboarding",
        "description": """• Warm welcome message: "Welcome to AriOme"
• Subtitle explaining the app's purpose
• Peaceful, minimalist design
• "Begin Your Journey" button to proceed"""
    },
    {
        "title": "3. Onboarding - Intentions",
        "route": "/onboarding (Step 2)",
        "description": """• Header: "What brings you here?"
• 8 selectable intention chips:
  - Healing (Pink), Growth (Green), Gratitude (Amber)
  - Presence (Purple), Trust (Blue), Creativity (Indigo)
  - Connection (Rose), Acceptance (Lime)
• Multi-select enabled with visual checkmarks
• "Continue to AriOme" button"""
    },
    {
        "title": "4. Authentication - Login",
        "route": "/auth",
        "description": """• "Welcome Back" title
• "Continue with Google" OAuth button
• Email and Password input fields
• "Sign In" button (teal accent)
• Toggle to switch to registration"""
    },
    {
        "title": "5. Authentication - Register",
        "route": "/auth (Register mode)",
        "description": """• "Create Account" title
• Google OAuth button
• Name, Email, Password fields
• "Create Account" button
• Toggle to switch to login"""
    },
    {
        "title": "6. Reflect Screen (Self-AriOme)",
        "route": "/(tabs)/reflect",
        "description": """• Personalized greeting: "Welcome, [User Name]"
• Mood Check-in: 10 mood options (Peaceful, Grateful, Hopeful, etc.)
• Today's Reflection: Daily prompt card
• Your Journey Stats: Reflections count, Day Streak"""
    },
    {
        "title": "7. Practices Screen",
        "route": "/(tabs)/practices",
        "description": """• Category Filters: All, Breathwork, Stillness, Gratitude, Body
• Practice Cards with title, duration, category, description
• Sample: Box Breathing, Morning Stillness, Gratitude Meditation
• Tap to open detailed practice view"""
    },
    {
        "title": "8. Wisdom Library",
        "route": "/(tabs)/wisdom",
        "description": """• Wisdom Cards with quotes and author attribution
• "Tap to reflect" prompt
• Resonance counter (heart icon)
• Reflect-after-consume flow with "This resonates" button"""
    },
    {
        "title": "9. Journal Screen",
        "route": "/(tabs)/journal",
        "description": """• "Your Journal" header
• List of past reflections with dates
• "+" button to create new entry
• Voice recording with OpenAI Whisper transcription
• Requires authentication for access"""
    },
    {
        "title": "10. Circles Screen (Community)",
        "route": "/(tabs)/circles",
        "description": """• Intention Filter Chips for filtering circles
• Circle Cards: name, description, member count, post count
• Sample Circles: Healing Hearts, Mindful Mornings, Growth Mindset
• Join/Leave functionality for authenticated users"""
    },
    {
        "title": "11. Profile Screen",
        "route": "/(tabs)/profile",
        "description": """• User avatar and name display
• Settings: Edit Profile, Intentions, Language, Notifications
• Privacy settings
• Logout button"""
    }
]

for screen in screens:
    doc.add_heading(screen["title"], level=1)
    doc.add_paragraph(f"Route: {screen['route']}")
    doc.add_paragraph(screen["description"])
    doc.add_paragraph('')

# Design System
doc.add_heading('Design System', level=1)
doc.add_heading('Colors', level=2)
doc.add_paragraph("""• Background Deep: #0A0A0F (Near black)
• Background Secondary: #1A1A24 (Dark gray)
• Primary Accent (Teal): #14B8A6
• Text Primary: #FFFFFF
• Text Muted: #9CA3AF""")

doc.add_heading('Navigation Structure', level=2)
doc.add_paragraph("""App
├── / (Welcome)
├── /onboarding
├── /auth
└── /(tabs)/
    ├── reflect (Home)
    ├── practices
    ├── wisdom
    ├── journal
    ├── circles
    └── profile""")

# Footer
doc.add_paragraph('')
doc.add_paragraph('Document generated: January 2025')
doc.add_paragraph('AriOme - Your Path to Inner Peace')

# Save
doc.save('/app/AriOme_Screens_Documentation.docx')
print("Document saved to /app/AriOme_Screens_Documentation.docx")
